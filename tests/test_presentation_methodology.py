# -*- coding: utf-8 -*-
"""
Tests for Slide 19 and 20 updates in AnekeshD_Progress.pptx.
Validates:
- Embedding of Figure 15 on Slide 19 with bilingual header and caption.
- Protection of native branding and navigation shapes on Slide 19.
- Updating of 4 milestone cards and header on Slide 20.
- Preservation of 16:9 widescreen dimensions and total slide count.
"""

import os
import unittest
from pptx import Presentation
from pptx.util import Inches


class TestPresentationMethodology(unittest.TestCase):
    def test_slide_19_and_20_updates(self):
        ppt_path = os.path.join(os.path.expanduser("~"), "Desktop", "AnekeshD_Progress.pptx")
        self.assertTrue(os.path.exists(ppt_path), f"{ppt_path} must exist")
        prs = Presentation(ppt_path)
        self.assertEqual(len(prs.slides), 23)

        # Widescreen 16:9 check
        self.assertAlmostEqual(prs.slide_width.inches, 13.333, places=2)
        self.assertAlmostEqual(prs.slide_height.inches, 7.500, places=2)

        # Slide 19: Check for embedded picture Figure 15
        s19 = prs.slides[18]
        pictures = [s for s in s19.shapes if s.shape_type == 13]
        self.assertGreater(len(pictures), 0, "Slide 19 must contain embedded Figure 15 picture")
        
        # Check picture dimensions
        pic = pictures[0]
        self.assertAlmostEqual(pic.left.inches, 0.60, delta=0.05)
        self.assertAlmostEqual(pic.top.inches, 1.85, delta=0.05)
        self.assertAlmostEqual(pic.width.inches, 12.133, delta=0.05)
        self.assertAlmostEqual(pic.height.inches, 4.85, delta=0.05)

        s19_text = " ".join(s.text_frame.text for s in s19.shapes if s.has_text_frame)
        self.assertIn("Methodological Innovations", s19_text)
        self.assertIn("面向低资源哈萨克语的AI生成文本检测与事实核验总体方法创新架构与技术流程", s19_text)
        self.assertIn("Figure 15", s19_text)

        # Slide 19: Check protected shapes
        s19_names = [s.name for s in s19.shapes]
        self.assertIn("灯片编号占位符 10", s19_names)
        self.assertIn("直接连接符 6", s19_names)
        self.assertIn("矩形 4", s19_names)
        self.assertIn("矩形 29", s19_names)

        # Slide 20: Check for header and milestone details
        s20 = prs.slides[19]
        s20_text = " ".join(s.text_frame.text for s in s20.shapes if s.has_text_frame)
        self.assertIn("Current Progress: LNCS Acceptance & September Meeting Agenda", s20_text)
        self.assertIn("论文录用进展、9月课题组汇报演示计划与后续工作推进安排", s20_text)
        self.assertIn("Springer LNCS", s20_text)
        self.assertIn("Gradio", s20_text)
        self.assertIn("September", s20_text)
        self.assertIn("Professor Guo", s20_text)


if __name__ == "__main__":
    unittest.main()
