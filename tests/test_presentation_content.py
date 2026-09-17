# -*- coding: utf-8 -*-
"""
Unit tests for content population, comparison tables, and figure embedding across all 23 slides.
Verifies:
1. All 14 figure slides (3, 4, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18) contain embedded pictures.
2. Formatted comparison tables exist on Slides 5, 11, 12, 14, 16 (and 22).
3. Slide 22 contains Committee comments with Reviewer 1, Reviewer 2, and '100.00% AUC'.
4. Captions (Figure X) and titles are present.
5. Zero emojis across all slides.
"""

import os
import unittest
from pptx import Presentation

OUTPUT_PPT_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "Kazakh_AI_Detection_Thesis_Progress_Prof_Guo.pptx"
)


class TestPresentationContent(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.prs_path = OUTPUT_PPT_PATH
        cls.assertTrue(
            os.path.isfile(cls.prs_path),
            f"Presentation file does not exist at {cls.prs_path}"
        )
        cls.prs = Presentation(cls.prs_path)

    def test_figures_and_tables_embedded_correctly(self):
        # Verify slides that must have embedded pictures
        fig_slides = [3, 4, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18]
        for s_num in fig_slides:
            slide = self.prs.slides[s_num - 1]
            pictures = [s for s in slide.shapes if s.shape_type == 13]  # MSO_SHAPE_TYPE.PICTURE
            self.assertGreater(
                len(pictures), 0,
                f"Slide {s_num} must have at least one embedded picture figure"
            )

        # Verify slides that must have tables
        table_slides = [5, 11, 12, 14, 16]
        for s_num in table_slides:
            slide = self.prs.slides[s_num - 1]
            tables = [s for s in slide.shapes if s.has_table]
            self.assertGreater(
                len(tables), 0,
                f"Slide {s_num} must contain an academic table"
            )

        # Verify Slide 22 Committee comments
        s22 = self.prs.slides[21]
        text_s22 = "".join(s.text_frame.text for s in s22.shapes if s.has_text_frame)
        for s in s22.shapes:
            if s.has_table:
                for row in s.table.rows:
                    for cell in row.cells:
                        text_s22 += cell.text
        self.assertIn("Reviewer 1", text_s22)
        self.assertIn("Reviewer 2", text_s22)
        self.assertIn("100.00% AUC", text_s22)

    def test_figure_captions_present(self):
        fig_slides = [3, 4, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18]
        for s_num in fig_slides:
            slide = self.prs.slides[s_num - 1]
            all_text = "".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
            self.assertIn(
                "Figure ", all_text,
                f"Slide {s_num} must contain a formal figure caption starting with 'Figure '"
            )

    def test_zero_emojis_across_deck(self):
        # Scan all slides for emojis
        for idx, slide in enumerate(self.prs.slides):
            for shp in slide.shapes:
                if shp.has_text_frame:
                    for ch in shp.text_frame.text:
                        code = ord(ch)
                        is_emoji = (
                            (0x1F600 <= code <= 0x1F64F) or
                            (0x1F300 <= code <= 0x1F5FF) or
                            (0x1F680 <= code <= 0x1F6FF) or
                            (0x1F700 <= code <= 0x1F77F) or
                            (0x1F900 <= code <= 0x1F9FF) or
                            (0x1FA00 <= code <= 0x1FA6F) or
                            (0x1FA70 <= code <= 0x1FAFF) or
                            (0x2600 <= code <= 0x26FF) or
                            (0x2700 <= code <= 0x27BF and code not in (0x2713, 0x2714))
                        )
                        self.assertFalse(
                            is_emoji,
                            f"Slide {idx+1} contains decorative emoji: '{ch}' (U+{code:04X})"
                        )

    def test_slide_07_full_width_methodological_framework(self):
        s7 = self.prs.slides[6]
        pictures = [s for s in s7.shapes if s.shape_type == 13]
        self.assertGreater(len(pictures), 0, "Slide 7 must contain Figure 3 picture")
        pic = pictures[0]
        self.assertAlmostEqual(pic.left.inches, 0.60, delta=0.05)
        self.assertAlmostEqual(pic.top.inches, 1.85, delta=0.05)
        self.assertAlmostEqual(pic.width.inches, 12.133, delta=0.05)
        self.assertAlmostEqual(pic.height.inches, 4.85, delta=0.05)

        s7_text = " ".join(s.text_frame.text for s in s7.shapes if s.has_text_frame)
        self.assertIn("Comprehensive Methodological Framework", s7_text)
        self.assertIn("Figure 3", s7_text)

    def test_slide_18_cards_and_zero_emoji_absence(self):
        s18 = self.prs.slides[17]
        pictures = [s for s in s18.shapes if s.shape_type == 13]
        self.assertGreater(len(pictures), 0, "Slide 18 must contain Figure 14 picture")

        s18_text = " ".join(s.text_frame.text for s in s18.shapes if s.has_text_frame)
        self.assertIn("314 / 314", s18_text)
        self.assertIn("Automated Test Suite", s18_text)
        self.assertIn("3 Formats", s18_text)
        self.assertIn("Multi-Format Ingestion", s18_text)
        self.assertIn("314 passing tests", s18_text)
        self.assertNotIn("Zero Emoji", s18_text)
        self.assertNotIn("zero emoji", s18_text.lower())

    def test_slide_01_presenter_notes_cleanup(self):
        s1 = self.prs.slides[0]
        self.assertTrue(s1.has_notes_slide)
        notes = s1.notes_slide.notes_text_frame.text
        self.assertNotIn("范千悦", notes)
        self.assertNotIn("陈亚兴", notes)
        self.assertNotIn("Fan Qianyue", notes)
        self.assertNotIn("Chen Yaxing", notes)
        self.assertIn("Daulet", notes)

    def test_slide_02_sequential_contents_ordering(self):
        s2 = self.prs.slides[1]
        order = []
        for s in s2.shapes:
            num = None
            if s.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                for sub in s.shapes:
                    if sub.has_text_frame and sub.text_frame.text.strip()[:2].isdigit():
                        num = int(sub.text_frame.text.strip()[:2])
            elif s.has_text_frame and s.text_frame.text.strip()[:2].isdigit():
                num = int(s.text_frame.text.strip()[:2])
            if num is not None:
                order.append(num)
        self.assertEqual(order, [1, 2, 3, 4, 5, 6], f"Slide 2 chapter iteration order must be [1, 2, 3, 4, 5, 6], got {order}")

    def test_zero_banned_terms_across_deck(self):
        banned = [
            "cross-attention", "cross attention", "cross-gate", "fan qianyue", "chen yaxing",
            "breakthrough", "catastrophic", "eliminates domain collapse", "production-ready",
            "perfect separation", "perfect auc", "311", "288"
        ]
        for idx, slide in enumerate(self.prs.slides):
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_lower = slide.notes_slide.notes_text_frame.text.lower()
                for b in banned:
                    self.assertNotIn(b, notes_lower, f"Slide {idx+1} notes contain banned term '{b}'")
            for s in slide.shapes:
                if s.has_text_frame:
                    txt_lower = s.text_frame.text.lower()
                    for b in banned:
                        self.assertNotIn(b, txt_lower, f"Slide {idx+1} shape {s.name} contains banned term '{b}'")
                if s.has_table:
                    for row in s.table.rows:
                        for cell in row.cells:
                            cell_lower = cell.text.lower()
                            for b in banned:
                                self.assertNotIn(b, cell_lower, f"Slide {idx+1} table cell contains banned term '{b}'")

    def test_slide_08_morphology_gated_fusion_terminology(self):
        s8 = self.prs.slides[7]
        s8_text = " ".join(s.text_frame.text for s in s8.shapes if s.has_text_frame)
        self.assertIn("Dual-Stream Morphology-Aware Gated Fusion", s8_text)
        self.assertNotIn("Cross-Attention", s8_text)

    def test_slide_09_robust_generalization_restoration(self):
        s9 = self.prs.slides[8]
        s9_text = " ".join(s.text_frame.text for s in s9.shapes if s.has_text_frame)
        self.assertIn("Robust Generalization & Adversarial Rewriting Defense", s9_text)
        self.assertIn("SentencePreservingChunker", s9_text)

    def test_slide_22_academic_status_badges(self):
        s22 = self.prs.slides[21]
        for s in s22.shapes:
            if s.has_table:
                for row in s.table.rows:
                    status = row.cells[3].text
                    self.assertNotIn("RESOLVED", status)
                    if status != "Status":
                        self.assertIn("Addressed in Thesis Chapter", status)

    def test_all_latin_fonts_are_times_new_roman(self):
        """Verify that all Latin text across the presentation uses Times New Roman and 0 Arial."""
        arial_found = []
        for s_idx, slide in enumerate(self.prs.slides, 1):
            for shp in slide.shapes:
                if shp.has_text_frame:
                    for p in shp.text_frame.paragraphs:
                        if p.font.name == "Arial":
                            arial_found.append(f"Slide {s_idx} para: {p.text[:30]}")
                        for r in p.runs:
                            if r.font.name == "Arial":
                                arial_found.append(f"Slide {s_idx} run: {r.text[:30]}")
                if shp.has_table:
                    for row in shp.table.rows:
                        for cell in row.cells:
                            for p in cell.text_frame.paragraphs:
                                if p.font.name == "Arial":
                                    arial_found.append(f"Slide {s_idx} table cell para: {p.text[:30]}")
                                for r in p.runs:
                                    if r.font.name == "Arial":
                                        arial_found.append(f"Slide {s_idx} table cell run: {r.text[:30]}")
        self.assertEqual(len(arial_found), 0, f"Found Arial font instances: {arial_found}")

        # Check theme fontScheme
        for rel in self.prs.part.rels.values():
            if "theme" in rel.target_ref:
                xml_text = rel.target_part.blob.decode("utf-8", errors="ignore")
                self.assertIn('<a:latin typeface="Times New Roman"/>', xml_text)
                self.assertNotIn('<a:latin typeface="Arial"/>', xml_text)

    def test_slide_06_hallucination_verification_related_work(self):
        """Verify Slide 6 contains structured LLM hallucination and 2D trust matrix related work."""
        s6 = self.prs.slides[5]
        s6_text = " ".join(s.text_frame.text for s in s6.shapes if s.has_text_frame)
        self.assertIn("LLM Hallucination Verification & Factual Grounding Gaps", s6_text)
        self.assertIn("The Scientific Blindspot: Detection Alone Cannot Verify Truth", s6_text)
        self.assertIn("1. Stylistic AI Detectors (Topics 1 & 2)", s6_text)
        self.assertIn("2. International Fact-Checking Corpora", s6_text)
        self.assertIn("3. Kazakh-FEVER & 2D Trust Matrix (Topic 3)", s6_text)
        self.assertNotIn("SciFact", s6_text)


if __name__ == "__main__":
    unittest.main()

