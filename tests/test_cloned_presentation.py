# -*- coding: utf-8 -*-
"""
Unit tests for cloned presentation template and branding engine.
Verifies:
1. Output presentation exists and has exactly 23 slides.
2. Native institutional branding is preserved on Slide 1 and content slides.
3. Presenter and advisor information is updated on Slide 1.
4. All commercial watermarks ('243001978', '合作QQ') are excised across all slides.
5. Content slides (3-22) retain layout '1_标题和内容'.
6. Gradient highlight tracker ('矩形 29') is correctly repositioned across sections.
7. Legacy poultry content shapes are cleared on content slides.
8. Slide 2 (TOC) and Slide 23 (Closing) are properly updated.
"""

import os
import unittest
from pptx import Presentation
from pptx.util import Inches

OUTPUT_PPT_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "Kazakh_AI_Detection_Thesis_Progress_Prof_Guo.pptx"
)


class TestClonedPresentation(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.prs_path = OUTPUT_PPT_PATH
        cls.file_exists = os.path.isfile(cls.prs_path)
        if cls.file_exists:
            cls.prs = Presentation(cls.prs_path)
        else:
            cls.prs = None

    def test_01_output_file_exists(self):
        self.assertTrue(
            self.file_exists,
            f"Expected output presentation does not exist at {self.prs_path}"
        )

    def test_02_slide_count_is_23(self):
        self.assertIsNotNone(self.prs, "Presentation could not be loaded")
        self.assertEqual(
            len(self.prs.slides),
            23,
            f"Expected exactly 23 slides, found {len(self.prs.slides)}"
        )

    def test_03_slide_1_branding_preserved(self):
        self.assertIsNotNone(self.prs, "Presentation could not be loaded")
        s1 = self.prs.slides[0]
        shape_names = [shp.name for shp in s1.shapes]

        # Verify KazNU logo, NPU logo group, and NPU background watermark
        self.assertIn("Изображение 16", shape_names, "Slide 1 missing KazNU logo (Изображение 16)")
        self.assertIn("LOGO组合", shape_names, "Slide 1 missing NPU logo group (LOGO组合)")
        self.assertIn("校徽打底", shape_names, "Slide 1 missing NPU watermark (校徽打底)")

    def test_04_slide_1_presenter_and_title_updated(self):
        self.assertIsNotNone(self.prs, "Presentation could not be loaded")
        s1 = self.prs.slides[0]
        all_text = ""
        for shp in s1.shapes:
            if shp.has_text_frame:
                all_text += " " + shp.text
            if shp.shape_type == 6:  # Group
                for sub in shp.shapes:
                    if sub.has_text_frame:
                        all_text += " " + sub.text

        self.assertIn("大雷 (Daulet)", all_text, "Slide 1 must contain presenter '大雷 (Daulet)'")
        self.assertIn("郭教授 (Prof. Guo)", all_text, "Slide 1 must contain advisor '郭教授 (Prof. Guo)'")
        self.assertIn(
            "Research on Morphologically-Grounded AI-Generated Text Detection",
            all_text,
            "Slide 1 must contain thesis title"
        )
        self.assertIn(
            "面向低资源哈萨克语的形态感知AI生成文本检测与事实核验研究",
            all_text,
            "Slide 1 must contain Chinese thesis subtitle"
        )

    def test_05_zero_commercial_watermarks(self):
        self.assertIsNotNone(self.prs, "Presentation could not be loaded")
        for i, slide in enumerate(self.prs.slides):
            for shp in slide.shapes:
                if shp.has_text_frame:
                    self.assertNotIn(
                        "243001978", shp.text,
                        f"Slide {i+1} shape '{shp.name}' contains commercial watermark '243001978'"
                    )
                    self.assertNotIn(
                        "合作QQ", shp.text,
                        f"Slide {i+1} shape '{shp.name}' contains commercial watermark '合作QQ'"
                    )
                if shp.shape_type == 6:  # Group
                    for sub in shp.shapes:
                        if sub.has_text_frame:
                            self.assertNotIn(
                                "243001978", sub.text,
                                f"Slide {i+1} group '{shp.name}' subshape '{sub.name}' contains '243001978'"
                            )
                            self.assertNotIn(
                                "合作QQ", sub.text,
                                f"Slide {i+1} group '{shp.name}' subshape '{sub.name}' contains '合作QQ'"
                            )

    def test_06_content_slides_layout_name(self):
        self.assertIsNotNone(self.prs, "Presentation could not be loaded")
        for i in range(2, 22):
            slide = self.prs.slides[i]
            layout_name = slide.slide_layout.name
            self.assertEqual(
                layout_name,
                "1_标题和内容",
                f"Slide {i+1} layout is '{layout_name}', expected '1_标题和内容'"
            )

    def test_07_gradient_tracker_repositioning(self):
        self.assertIsNotNone(self.prs, "Presentation could not be loaded")
        tolerance = Inches(0.08)

        expected_tracker_positions = [
            (2, Inches(1.30), Inches(1.82)),   # Slide 3 (Section 1 Background)
            (3, Inches(1.30), Inches(1.82)),   # Slide 4 (Section 1 Background)
            (4, Inches(3.12), Inches(1.66)),   # Slide 5 (Section 2 Related Work)
            (5, Inches(3.12), Inches(1.66)),   # Slide 6 (Section 2 Related Work)
            (6, Inches(4.78), Inches(1.81)),   # Slide 7 (Section 3 Content)
            (9, Inches(4.78), Inches(1.81)),   # Slide 10 (Section 3 Content)
            (10, Inches(6.59), Inches(2.15)),  # Slide 11 (Section 4 Experiments)
            (15, Inches(6.59), Inches(2.15)),  # Slide 16 (Section 4 Experiments)
            (16, Inches(6.59), Inches(2.15)),  # Slide 17 (Section 5 Engineering)
            (17, Inches(6.59), Inches(2.15)),  # Slide 18 (Section 5 Engineering)
            (18, Inches(8.92), Inches(1.95)),  # Slide 19 (Section 6 Conclusion)
            (20, Inches(8.92), Inches(1.95)),  # Slide 21 (Section 6 Conclusion)
            (21, Inches(10.83), Inches(2.00)), # Slide 22 (Comments and Responses)
        ]

        for s_idx, exp_left, exp_width in expected_tracker_positions:
            slide = self.prs.slides[s_idx]
            r29_list = [s for s in slide.shapes if s.name == "矩形 29"]
            self.assertEqual(
                len(r29_list), 1,
                f"Slide {s_idx+1} expected exactly one '矩形 29', found {len(r29_list)}"
            )
            r29 = r29_list[0]
            self.assertAlmostEqual(
                r29.left, exp_left, delta=tolerance,
                msg=f"Slide {s_idx+1} tracker left {r29.left} does not match expected {exp_left}"
            )
            self.assertAlmostEqual(
                r29.width, exp_width, delta=tolerance,
                msg=f"Slide {s_idx+1} tracker width {r29.width} does not match expected {exp_width}"
            )

    def test_08_legacy_shapes_cleared(self):
        self.assertIsNotNone(self.prs, "Presentation could not be loaded")
        legacy_poultry_terms = [
            "poultry", "chickens", "broiler", "chicken coop",
            "肉鸡", "养殖", "蛋鸡", "鸡舍", "雏鸡", "出栏", "屠宰", "饲料", "家禽"
        ]
        for i in range(2, 22):
            slide = self.prs.slides[i]
            all_text = " ".join(shp.text for shp in slide.shapes if shp.has_text_frame).lower()
            for term in legacy_poultry_terms:
                self.assertNotIn(
                    term.lower(), all_text,
                    f"Slide {i+1} still contains legacy poultry term '{term}'"
                )

    def test_09_slide_2_table_of_contents(self):
        self.assertIsNotNone(self.prs, "Presentation could not be loaded")
        s2 = self.prs.slides[1]
        all_text = ""
        for shp in s2.shapes:
            if shp.has_text_frame:
                all_text += " " + shp.text
            if shp.shape_type == 6:
                for sub in shp.shapes:
                    if sub.has_text_frame:
                        all_text += " " + sub.text

        self.assertIn("Research Background", all_text)
        self.assertIn("Related Work", all_text)
        self.assertIn("Research Content", all_text)
        self.assertIn("Experiments", all_text)
        self.assertIn("Engineering", all_text)
        self.assertIn("Conclusion", all_text)

    def test_10_slide_23_closing(self):
        self.assertIsNotNone(self.prs, "Presentation could not be loaded")
        s23 = self.prs.slides[22]
        all_text = ""
        for shp in s23.shapes:
            if shp.has_text_frame:
                all_text += " " + shp.text

        self.assertIn("大雷 (Daulet)", all_text)
        self.assertIn("郭教授 (Prof. Guo)", all_text)
        self.assertTrue(
            "Thank You for Your Attention!" in all_text or "Thank you for attention!" in all_text
        )


if __name__ == "__main__":
    unittest.main()
