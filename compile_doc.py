import re
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    # Helper to set cell padding in docx (in dxa: 20 dxa = 1 pt)
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        m = OxmlElement(f'w:{name}')
        m.set(qn('w:w'), str(val))
        m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    tcPr.append(tcMar)

def add_table_borders(table):
    # Helper to add standard grid borders to a table
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')  # 4 eighths of a pt = 0.5 pt
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'CCCCCC')
        borders.append(border)
    tblPr.append(borders)

def compile_markdown_to_docx(md_path, docx_path):
    doc = Document()
    
    # Page setup (margins)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Read Markdown
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_rows = []

    # Parse line-by-line
    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')
        
        # Skip divider lines in tables
        if in_table and re.match(r'^\s*\|?\s*[:\-]+\s*\|', line):
            i += 1
            continue

        # Check if table ends
        if in_table and (not line.strip() or not line.startswith('|')):
            # Render the collected table
            if table_rows:
                headers = [h.strip() for h in table_rows[0].split('|')[1:-1]]
                col_count = len(headers)
                table = doc.add_table(rows=len(table_rows), cols=col_count)
                table.autofit = True
                add_table_borders(table)
                
                # Populate table
                for r_idx, row_str in enumerate(table_rows):
                    cols = [c.strip() for c in row_str.split('|')[1:-1]]
                    # Pad if necessary
                    while len(cols) < col_count:
                        cols.append("")
                    for c_idx, val in enumerate(cols[:col_count]):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = val
                        set_cell_margins(cell)
                        # Formatting headers
                        if r_idx == 0:
                            for run in cell.paragraphs[0].runs:
                                run.font.bold = True
                                run.font.size = Pt(10.5)
                        else:
                            for run in cell.paragraphs[0].runs:
                                run.font.size = Pt(9.5)
            table_rows = []
            in_table = False
            if not line.strip():
                i += 1
                continue

        # Table start/row
        if line.startswith('|'):
            in_table = True
            table_rows.append(line)
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
        elif line.startswith('## '):
            p = doc.add_heading(line[3:], level=2)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
        elif line.startswith('### '):
            p = doc.add_heading(line[4:], level=3)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
        
        # List Items
        elif line.strip().startswith(('* ', '- ')):
            clean_line = line.strip()[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2.5)
            p.add_run(clean_line)
        elif re.match(r'^\s*\d+\.\s', line):
            # Parse as a normal paragraph with manual numbering to avoid Word's list continuation bug
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.15
            
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                else:
                    p.add_run(part)

        # Paragraph
        elif line.strip():
            # Check if it is a thematic break/horizontal rule
            if line.strip() == '---':
                i += 1
                continue
            
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            
            # Simple inline formatting parser (bold/italics)
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                else:
                    p.add_run(part)

        i += 1
        
    doc.save(docx_path)
    print(f"Successfully compiled {md_path} to {docx_path}")

if __name__ == '__main__':
    compile_markdown_to_docx('paper_draft.md', 'paper_draft.docx')
